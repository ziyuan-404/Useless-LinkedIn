"""Read the historical workbook. Never changes the source file."""
import datetime
import json
import re
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')
book = load_workbook(sys.argv[1], read_only=False, data_only=True)
records = []
for sheet in book:
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', sheet.title):
        continue
    headers = [sheet.cell(1, col).value for col in range(1, 24)]
    if headers[0] not in ('记录 ID', '記錄 ID') or headers[13] not in ('岗位链接', '职位链接', '岗位URL', 'Job URL'):
        raise ValueError(f'Unexpected headers in {sheet.title}')
    for cells in sheet.iter_rows(min_row=2, max_col=23):
        if not any(cell.value is not None for cell in cells):
            continue
        values = []
        for cell in cells:
            value = cell.value
            if cell.column == 14 and cell.hyperlink and cell.hyperlink.target:
                value = cell.hyperlink.target
            if isinstance(value, (datetime.date, datetime.datetime)):
                value = value.isoformat()[:10]
            values.append('' if value is None else str(value))
        records.append({'sheet': sheet.title, 'row': cells[0].row, 'values': values})
print(json.dumps(records, ensure_ascii=False))

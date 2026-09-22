"""Read posting links from dated dashboard sheets without modifying the workbook."""
import json
import re
import sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding='utf-8')
workbook = load_workbook(sys.argv[1], read_only=False, data_only=True)
records = []
for sheet in workbook:
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', sheet.title):
        continue
    column = None
    header_row = None
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10)):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() in ('岗位链接', '职位链接', '岗位URL', 'Job URL'):
                column, header_row = cell.column, cell.row
                break
        if column is not None:
            break
    if column is None:
        continue
    for row in sheet.iter_rows(min_row=header_row + 1, min_col=column, max_col=column):
        cell = row[0]
        url = (cell.hyperlink.target or cell.value) if cell.hyperlink else cell.value
        if isinstance(url, str) and re.match(r'^https?://', url.strip()):
            records.append({'url': url.strip(), 'sheet': sheet.title, 'row': cell.row})
print(json.dumps(records, ensure_ascii=False))
